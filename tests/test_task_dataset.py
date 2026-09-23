"""One row per iteration, flexible variables, and a compact, lossless export."""
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from kildeanalyse import tjeneste
from kildeanalyse.adaptere import ADAPTERE
from kildeanalyse.modell import Motorsvar, Plan
from kildeanalyse.task_dataset import describe, add_result, plan_preview
from test_tasks import fixture, QUOTE


def obj(properties):
    return {'type': 'object', 'properties': properties, 'required': list(properties), 'additionalProperties': False}


def execute(tmp_path, monkeypatch, schema, values, *, limitations=()):
    store, aid, _, _ = fixture(tmp_path, monkeypatch)
    tjeneste.ny_planversjon(store, aid, 'Dataset contract', output_schema=schema)
    runs = tjeneste.legg_til_kjoringer(store, aid)['nye']
    def read(self, package, *args):
        result = values[0 if package.dokument_navn == 'a.txt' else 1]
        response = {'result': result, 'source_units_read': [s.nr for s in package.sider], 'limitations': list(limitations)}
        return Motorsvar(json.dumps(response), response, modell_rapportert='fixture')
    monkeypatch.setattr(ADAPTERE['claude_cli'], 'kjor', read)
    tjeneste.godkjenn_plan(store, aid, 'Test fixture')
    tjeneste.start(store, aid, kjoring_ider=[r['id'] for r in runs])
    return store, aid, runs


def rows(ws):
    values = list(ws.values)
    return [dict(zip(values[0], r)) for r in values[1:]]


def test_object_variables_keep_types_and_one_row_per_run(tmp_path, monkeypatch):
    schema = obj({'entity': {'type': 'string', 'title': 'Organisation', 'description': 'Named organisation.'},
                  'metrics': obj({'score': {'type': ['number', 'null'], 'description': 'Score on a 0–10 scale.'},
                                  'count': {'type': 'integer'}, 'positive': {'type': 'boolean'}})})
    store, aid, runs = execute(tmp_path, monkeypatch, schema,
        [{'entity': 'A', 'metrics': {'score': 0, 'count': 0, 'positive': False}},
         {'entity': 'B', 'metrics': {'score': None, 'count': 3, 'positive': True}}])
    preview = tjeneste.vis_plan(store, aid)['dataset_preview']
    assert preview['structured'] and 'One row per run' in preview['row_unit']
    out = tjeneste.eksporter(store, aid, include_csv=True)
    book = load_workbook(out['workbook'])
    result = rows(book['Results'])
    assert len(result) == len(store.kjoringer(aid))  # Includes earlier unstarted plan-version runs.
    active = [r for r in result if r['Run'] in {x['id'] for x in runs}]
    keys = list(active[0])
    score = next(k for k in keys if 'metrics.score' in k)
    positive = next(k for k in keys if 'metrics.positive' in k)
    assert [r[score] for r in active] == [0, None]
    assert [r[positive] for r in active] == [False, True]
    assert book.sheetnames[0] == 'Results'
    assert 'Score on a 0–10 scale.' in str(list(book['Variables'].values))
    assert '/metrics/score' in str(list(book['Runs'].values))
    with zipfile.ZipFile(out['documentation_archive']) as archive:
        assert 'datasets/Results.csv' in archive.namelist()
        assert json.loads(archive.read(f'results/{runs[1]["id"]}.json'))['metrics']['score'] is None
    assert len(list(Path(out['mappe']).iterdir())) == 3


def test_nested_arrays_never_multiply_main_rows_or_cross_join(tmp_path, monkeypatch):
    entity = obj({'name': {'type': 'string'}, 'mentions': {'type': 'array', 'items': {'type': 'integer'}}})
    schema = obj({'score': {'type': 'number'}, 'entities': {'type': 'array', 'items': entity},
                  'topics': {'type': 'array', 'items': {'type': 'string'}}})
    store, aid, runs = execute(tmp_path, monkeypatch, schema,
        [{'score': 2.5, 'entities': [{'name':'X','mentions':[1,2]}, {'name':'Y','mentions':[3]}], 'topics':['a','b','c']},
         {'score': 0, 'entities':[], 'topics':[]}], limitations=['Limited evidence.'])
    out = tjeneste.eksporter(store, aid)
    book = load_workbook(out['workbook'])
    assert book['Results'].max_row - 1 == len(store.kjoringer(aid))
    details = {s['sheet']:s['rows'] for s in out['datasets'][1:]}
    assert sorted(details.values()) == [2,3,3]
    assert not any(c.value == '1. X\n\n2. Y' for row in book['Results'] for c in row)
    assert not any('entities.name' in str(c.value) for c in book['Results'][1])
    inline = load_workbook(tjeneste.eksporter(store, aid, list_layout='inline')['workbook'])
    assert any(c.value == '1. X\n\n2. Y' for row in inline['Results'] for c in row)
    assert len(rows(book['entities'])) == 2
    assert len(rows(book['entities.mentions'])) == 3
    parents = {r['Record ID'] for r in rows(book['entities'])}
    assert all(r['Parent record ID'] in parents for r in rows(book['entities.mentions']))
    issues = rows(book['Errors and notes'])
    assert len(issues) == 2
    assert all(r['Message'] == 'Limited evidence.' for r in issues)
    assert "['Limited evidence.']" not in str(list(book['Runs'].values))


def test_root_list_quotes_have_main_row_and_source_located_detail(tmp_path, monkeypatch):
    store, aid, runs, _ = fixture(tmp_path, monkeypatch, structured=True)
    tjeneste.godkjenn_plan(store, aid, 'Test fixture')
    tjeneste.start(store, aid)
    out = tjeneste.eksporter(store, aid)
    book = load_workbook(out['workbook'])
    assert book['Results'].max_row == 3
    assert len(rows(book['Findings'])) == 2
    assert all(r['Source location'] == 'Line 2' for r in rows(book['Findings']))
    assert all(r['Result items (count)'] == 1 for r in rows(book['Results']))
    assert 'passage' not in rows(book['Results'])[0]


def test_failure_and_rejection_retain_blank_main_rows(tmp_path, monkeypatch):
    store, aid, runs, _ = fixture(tmp_path, monkeypatch, broken=True)
    tjeneste.godkjenn_plan(store, aid, 'Test fixture')
    tjeneste.start(store, aid)
    attempt = store.forsok_for_kjoring(runs[1]['id'])[0]
    tjeneste.registrer_kontroll(store, attempt['id'], 'Test reviewer', 'avvist', 'Fixture rejection')
    out = tjeneste.eksporter(store, aid)
    book = load_workbook(out['workbook'])
    assert len(rows(book['Results'])) == 2
    assert all(r['Result'] is None for r in rows(book['Results']))
    assert 'Errors and notes' in book.sheetnames
    assert 'Rejected' in str(list(book['Runs'].values))
    with zipfile.ZipFile(out['documentation_archive']) as archive:
        assert archive.read(f'audit/attempts/{attempt["id"]}/raasvar.txt').decode('utf-8') == attempt['raasvar']


def test_formula_like_and_long_values_survive_inside_workbook(tmp_path, monkeypatch):
    long = 'A very long source passage. ' * 1800
    schema = obj({'text': {'type':'string'}, 'identifier': {'type':'integer'}})
    store, aid, runs = execute(tmp_path, monkeypatch, schema,
        [{'text': '=HYPERLINK("https://invalid.example")', 'identifier': 12345678901234567890},
         {'text': long, 'identifier': 0}])
    out = tjeneste.eksporter(store, aid)
    book = load_workbook(out['workbook'])
    cell = next(c for row in book['Results'] for c in row if isinstance(c.value,str) and c.value.startswith('=HYPERLINK'))
    assert cell.data_type == 's'
    assert '12345678901234567890' in str(list(book['Results'].values))
    parts = rows(book['Long texts'])
    assert ''.join(r['Text'] for r in parts) == long
    assert all(len(r['Text']) <= 2000 for r in parts)


def test_workbook_failure_does_not_publish_snapshot(tmp_path, monkeypatch):
    store, aid, runs, _ = fixture(tmp_path, monkeypatch)
    def fail(*args, **kwargs):
        raise RuntimeError('fixture export failure')
    monkeypatch.setattr('kildeanalyse.task_workbook.write_workbook', fail)
    with pytest.raises(RuntimeError, match='fixture export failure'):
        tjeneste.eksporter(store, aid)
    project = store.prosjekt(store.analyse(aid)['prosjekt_id'])
    assert list((Path(project['directory'])/'exports').iterdir()) == []


def test_empty_root_lists_and_null_results_are_not_failed_or_negative(tmp_path, monkeypatch):
    store, aid, runs = execute(tmp_path, monkeypatch, {'type':['array','null'], 'items':{'type':'string'}}, [[], None])
    out = tjeneste.eksporter(store, aid)
    book = load_workbook(out['workbook'])
    assert book['Results'].max_row - 1 == len(store.kjoringer(aid))
    states = [r['Dataset'] for r in rows(book['Runs'])]
    assert 'Empty result' in states and 'Result is null' in states
    active = {r['Run']: r for r in rows(book['Results'])}
    count_key = next(k for k in active[runs[0]['id']] if 'Result items (count)' in k)
    assert active[runs[0]['id']][count_key] == 0
    assert active[runs[1]['id']][count_key] is None
    assert not rows(book['Findings'])  # null must not invent a detail record



def test_preview_handles_prose_without_inventing_variables():
    preview = plan_preview(Plan('Summarize',task_instructions='Summarize'))
    assert not preview['structured']
    assert len(preview['tables'][0]['variables']) == 1
    assert 'Prose-only' in preview['notes'][-1]


def test_nullable_collection_keeps_one_count_and_distinguishes_empty():
    tables = describe(obj({'items': {'type':['array','null'], 'items':{'type':'integer'}}}))
    add_result(tables, {'items':None}, {'attempt_id':'null'})
    add_result(tables, {'items':[]}, {'attempt_id':'empty'})
    assert list(tables[()].columns) == [(('items',), 'count')]
    assert [r['values'][(('items',),'count')] for r in tables[()].rows] == [None, 0]
    assert not tables[('items','*')].rows
