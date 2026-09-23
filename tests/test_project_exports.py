import json
import shutil
import sqlite3
from pathlib import Path
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

from kildeanalyse import tjeneste
from kildeanalyse.lager import Lager
from kildeanalyse.project_files import root_for

FIX = Path(__file__).parent/'fixtures'/'syntetisk'


def setup(store, tmp_path, language='nb', settings=None):
    project = tjeneste.opprett_prosjekt(store, 'Policy #1 æøå', str(tmp_path/'Arbeidsmappe æøå'))
    tjeneste.importer_dokumenter(store, project['id'], [str(FIX/'fjordblikk_2025.pdf')])
    created = tjeneste.opprett_analyse(store, project['id'], 'Eksempel', 'Les rapporten', motor='simulert', sprak=language, motorinnstillinger=settings)
    aid = created['analyse']['id']
    tjeneste.godkjenn_plan(store, aid, 'Testperson')
    run = tjeneste.legg_til_kjoringer(store, aid)['nye'][0]
    return project, aid, run





def test_project_path_guards(tmp_path, monkeypatch):
    data = tmp_path/'data'
    store = Lager(data)
    project = store.opprett_prosjekt('Original')
    project_id = project['id']
    chosen = tmp_path/'Chosen'
    result = tjeneste.set_project_directory(store, project_id, str(chosen))
    assert result['directory'] == str(chosen)
    assert Lager(data).prosjekt(project_id)['directory'] == str(chosen)
    other = store.opprett_prosjekt('Other')
    with pytest.raises(ValueError, match='another project'):
        tjeneste.set_project_directory(store, other['id'], str(chosen))
    with pytest.raises(ValueError, match='absolute'):
        tjeneste.set_project_directory(store, other['id'], 'relative')
    with pytest.raises(ValueError, match='AppData'):
        tjeneste.set_project_directory(store, other['id'], str(data/'bad'))
    existing = tmp_path/'not-empty'
    existing.mkdir(); (existing/'keep.txt').write_text('user work')
    with pytest.raises(ValueError, match='empty'):
        tjeneste.set_project_directory(store, other['id'], str(existing))
    assert (existing/'keep.txt').read_text() == 'user work'
    changed = tjeneste.set_project_directory(store, project_id, str(tmp_path/'New'))
    assert chosen.exists() and Path(changed['directory']).exists()


def test_failed_run_is_visible_and_failed_export_does_not_publish(tmp_path, monkeypatch):
    store = Lager(tmp_path/'data')
    project, aid, run = setup(store, tmp_path, settings={'scenarier':{'fjordblikk_2025.pdf':'krasj'}})
    tjeneste.start(store, aid)
    output = tjeneste.eksporter(store, aid, med_kilder=False)
    book = load_workbook(output['workbook'])
    assert book['Resultater'].max_row == 2
    assert 'Feil og merknader' in book.sheetnames
    before = set((Path(project['directory'])/'exports').iterdir())
    def fail(*args, **kwargs): raise OSError('copy failure')
    monkeypatch.setattr('kildeanalyse.task_export.copy_file', fail)
    with pytest.raises(OSError, match='copy failure'):
        tjeneste.eksporter(store, aid)
    assert set((Path(project['directory'])/'exports').iterdir()) == before
