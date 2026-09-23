"""Reader exports stay compact across retries without hiding missing results."""
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from kildeanalyse import tjeneste
from test_tasks import fixture
from test_task_dataset import rows


def test_latest_failed_run_does_not_fall_back_and_full_history_survives(tmp_path, monkeypatch):
    store, aid, old, _ = fixture(tmp_path, monkeypatch, structured=True)
    tjeneste.godkjenn_plan(store, aid, 'Test')
    tjeneste.start(store, aid)
    # Cross numeric run-ID boundary and keep the same schema across versions.
    for _ in range(5):
        tjeneste.ny_planversjon(store, aid, 'Same task, changed reader settings')
        latest = tjeneste.legg_til_kjoringer(store, aid)['nye']
    out = tjeneste.eksporter(store, aid, include_csv=True)
    book = load_workbook(out['workbook'])
    main = rows(book['Results'])
    assert len(main) == 2
    assert {r['Run'] for r in main} == {r['id'] for r in latest}
    assert all('planlagt' in r['Status'] for r in main)
    assert all(r['Result items (count)'] is None for r in main)
    assert book.sheetnames == ['Results']  # no empty findings, definitions or telemetry tabs
    with zipfile.ZipFile(out['documentation_archive']) as z:
        assert f'results/{old[0]["id"]}.json' in z.namelist()
        assert 'datasets/variables.json' in z.namelist()
        assert len(json.loads(z.read('audit/analysis.json'))['runs']) == 12
    historical = tjeneste.eksporter(store, aid, row_scope='runs')
    assert load_workbook(historical['workbook'])['Results'].max_row == 13


def test_same_schema_across_selected_versions_shares_columns_and_details(tmp_path, monkeypatch):
    store, aid, original, _ = fixture(tmp_path, monkeypatch, structured=True)
    tjeneste.godkjenn_plan(store, aid, 'Test')
    tjeneste.start(store, aid)
    tjeneste.ny_planversjon(store, aid, 'Same schema')
    new = tjeneste.legg_til_kjoringer(store, aid, [original[0]['dokument_id']])['nye']
    tjeneste.godkjenn_plan(store, aid, 'Test')
    tjeneste.start(store, aid, kjoring_ider=[r['id'] for r in new])
    out = tjeneste.eksporter(store, aid)
    book = load_workbook(out['workbook'])
    assert book.sheetnames == ['Results', 'Findings']
    assert book['Results'].max_column == 5
    assert [r['Result items (count)'] for r in rows(book['Results'])] == [1, 1]
    assert len(rows(book['Findings'])) == 2
    assert all('Record ID' in r for r in rows(book['Findings']))


def test_latest_invalid_result_has_visible_status_and_reason(tmp_path, monkeypatch):
    store, aid, _, _ = fixture(tmp_path, monkeypatch, broken=True)
    tjeneste.godkjenn_plan(store, aid, 'Test')
    tjeneste.start(store, aid)
    out = tjeneste.eksporter(store, aid)
    book = load_workbook(out['workbook'])
    by_doc = {r['Document']: r for r in rows(book['Results'])}
    assert 'valideringsfeil' in by_doc['a.txt']['Status']
    assert by_doc['a.txt']['Result'] is None
    assert 'a.txt' in str(list(book['Errors and notes'].values))
    assert by_doc['b.txt']['Result']


def test_changed_definitions_are_not_merged(tmp_path, monkeypatch):
    store, aid, runs, _ = fixture(tmp_path, monkeypatch)
    tjeneste.ny_planversjon(store, aid, 'Different definition', output_schema={'type':'string', 'description':'New meaning'})
    tjeneste.legg_til_kjoringer(store, aid, [runs[0]['dokument_id']])
    book = load_workbook(tjeneste.eksporter(store, aid)['workbook'])
    assert book['Results'].max_column == 6
    assert len([c.value for c in book['Results'][1] if ': Result' in str(c.value)]) == 2


def test_short_destination_preserves_project_and_rejects_long_excel_path(tmp_path, monkeypatch):
    store, aid, _, _ = fixture(tmp_path, monkeypatch)
    project_id = store.analyse(aid)['prosjekt_id']
    before = store.prosjekt(project_id)['directory']
    parent = tmp_path / 'export'
    out = tjeneste.eksporter(store, aid, output_directory=str(parent))
    assert Path(out['workbook']).is_relative_to(parent)
    assert store.prosjekt(project_id)['directory'] == before
    assert len(Path(out['mappe']).name) < 20
    import os
    if os.name == 'nt':
        long = tmp_path / ('x'*90) / ('y'*90)
        with pytest.raises((ValueError, tjeneste.TjenesteFeil), match='output_directory'):
            tjeneste.eksporter(store, aid, output_directory=str(long))
        assert not long.exists()


def test_invalid_row_scope_is_rejected_before_publication(tmp_path, monkeypatch):
    store, aid, _, _ = fixture(tmp_path, monkeypatch)
    with pytest.raises((ValueError, tjeneste.TjenesteFeil), match='row_scope'):
        tjeneste.eksporter(store, aid, row_scope='best_success')
