import csv
import json
import sqlite3
from pathlib import Path
import pytest
from docx import Document
from openpyxl import Workbook
from kildeanalyse import tjeneste, visning
from kildeanalyse.dokument import importer_dokument, sider_uten_tekst, DokumentFeil, sha256_fil
from kildeanalyse.lager import Lager, SKJEMA, KJ_FULLFORT
from kildeanalyse.modell import Motorsvar, Stotte
from kildeanalyse.adaptere.claude_cli import ClaudeCliAdapter
from kildeanalyse.source_formats import extract, location
from kildeanalyse.task_contract import validate as valider

QUOTE = 'The board adopted a policy.'


def source_file(tmp_path, extension):
    path = tmp_path/f'example.{extension}'
    if extension == 'docx':
        doc = Document(); doc.add_heading('Policy', 1); doc.add_paragraph(QUOTE)
        table = doc.add_table(rows=1, cols=1); table.cell(0,0).text = 'Follow-up'
        doc.save(path)
    elif extension == 'xlsx':
        book = Workbook(); sheet = book.active; sheet.title = 'Governance'
        sheet['A1'] = 'Policy'; sheet['B2'] = QUOTE
        book.save(path); book.close()
    elif extension in ('csv','tsv'):
        with path.open('w',encoding='utf-8-sig',newline='') as f:
            writer = csv.writer(f,delimiter=';' if extension == 'csv' else '\t')
            writer.writerows([['Name','Statement'], ['Example',QUOTE]])
    else:
        path.write_text('Policy\n\n'+QUOTE,encoding='utf-8')
    return path


@pytest.mark.parametrize('extension,locator', [('txt','Line 3'),('md','Line 3'),('csv','Record 2'),
                                               ('tsv','Record 2'),('docx','Body/block 2'),('xlsx','Governance!B2:B2')])
def test_complete_format_workflow_without_models(tmp_path, monkeypatch, extension, locator):
    file = source_file(tmp_path, extension); original_hash = sha256_fil(file)
    store = Lager(tmp_path/'data'); project = tjeneste.opprett_prosjekt(store, 'Comparable files')
    document, new = importer_dokument(store, project['id'], file)
    assert new and not sider_uten_tekst(document)
    assert importer_dokument(store, project['id'], file)[1] is False
    assert Path(document['lagret_kopi']).suffix == '.'+extension
    criterion = {'criteria':[{'id':'policy', 'question':'Does the source describe a policy?',
                              'allowed_answers':['yes','not_mentioned'], 'evidence_required_for':['yes']}]}
    analysis = tjeneste.opprett_analyse(store, project['id'], 'Policy', 'Compare each file', motor='claude_cli', sprak='en')
    aid = analysis['analyse']['id']; plan = analysis['planversjon']['plan']
    assert tjeneste.vis_plan(store,aid)['source_profiles'][0]['format'] == extension
    run = tjeneste.legg_til_kjoringer(store,aid)['nye'][0]['id']
    preview = tjeneste.vis_inputpakke(store,run)['pakke']
    assert locator in preview['brukermelding'] and '[Source unit' in preview['brukermelding']
    assert 'fysiske sider' not in preview['brukermelding']
    monkeypatch.setattr(ClaudeCliAdapter, 'sjekk_stotte', lambda self:Stotte(True))
    def reader(self, package, *args):
        unit = next(s for s in package.sider if QUOTE in s.tekst)
        answer = {'result':QUOTE, 'source_units_read':[s.nr for s in package.sider], 'limitations':[]}
        return Motorsvar(raasvar=json.dumps(answer), svar=answer)
    monkeypatch.setattr(ClaudeCliAdapter, 'kjor', reader)
    tjeneste.godkjenn_plan(store,aid,'Offline check'); tjeneste.start(store,aid)
    assert store.kjoring(run)['status'] == KJ_FULLFORT
    details = tjeneste.vis_kjoring(store,run)
    assert details['forsok'][0]['result'] == QUOTE
    import zipfile
    output = tjeneste.eksporter(store, aid)
    with zipfile.ZipFile(output['documentation_archive']) as archive:
        assert json.loads(archive.read(f'results/{run}.json')) == QUOTE
        assert any(n.startswith('sources/') and n.endswith('.'+extension) for n in archive.namelist())
        audit = json.loads(archive.read(f'audit/{run}.json'))
        assert audit['dokument']['source_metadata']['format'] == extension
    assert sha256_fil(file) == original_hash


def test_workbook_formulas_hidden_sheet_and_short_values(tmp_path):
    file = tmp_path/'formulas.xlsx'; book = Workbook(); sheet = book.active
    sheet.title = 'Data'; sheet['B1'] = 17; sheet['C1'] = '=B1*2'
    hidden = book.create_sheet('Hidden'); hidden.sheet_state='hidden'; hidden['A3']='yes'
    book.save(file); book.close()
    units, profile, _ = extract(file)
    assert profile['missing_formula_values'] == ['Data!C1']
    assert 'UNAVAILABLE' in units[0]['tekst'] and '=B1*2' in units[0]['tekst']
    assert units[1]['source']['sheet_state'] == 'hidden' and units[1]['source']['row'] == 3


def test_csv_record_ids_handle_multiline_quotes_and_single_columns(tmp_path):
    file=tmp_path/'table.csv'; file.write_text('name;statement\nX;"first\nsecond"\n',encoding='utf-8')
    units, _, _ = extract(file)
    assert len(units) == 2 and units[1]['source']['record'] == 2 and 'first\nsecond' in units[1]['tekst']
    file.write_text('header\nvalue\n',encoding='utf-8')
    assert len(extract(file)[0]) == 2


def test_explicit_unsupported_and_empty_files_fail_visibly(tmp_path):
    store=Lager(tmp_path/'data'); project=store.opprett_prosjekt('Files')
    for extension, data in [('doc',b'unsupported'),('txt',b''),('txt',b'\xffbad encoding')]:
        file=tmp_path/f'bad.{extension}'; file.write_bytes(data)
        with pytest.raises(DokumentFeil): importer_dokument(store,project['id'],file)
    assert store.dokumenter(project['id']) == []


def test_folder_import_reports_exclusions_failures_and_duplicates(tmp_path):
    from kildeanalyse.languages import public_result
    store = Lager(tmp_path/'data'); project = store.opprett_prosjekt('Scope')
    sources = tmp_path/'sources'; sources.mkdir()
    good = sources/'report.TXT'; good.write_text(QUOTE, encoding='utf-8')
    empty = sources/'empty.txt'; empty.write_text('', encoding='utf-8')
    unsupported = sources/'slides.pptx'; unsupported.write_bytes(b'not supported')
    nested = sources/'earlier-years'; nested.mkdir()
    (nested/'older.txt').write_text('Older annual report.', encoding='utf-8')
    before = {p: sha256_fil(p) for p in sources.rglob('*') if p.is_file()}

    result = tjeneste.importer_dokumenter(store, project['id'], [str(sources)])
    rows = {Path(r['sti']).name: r for r in result['resultater']}
    assert set(rows) == {'report.TXT', 'empty.txt'}
    assert rows['report.TXT']['nytt'] is True and 'feil' in rows['empty.txt']
    assert {r['path']: r['reason'] for r in result['skipped']} == {
        str(unsupported): 'unsupported_format', str(nested): 'subdirectory'}
    # Both MCP presentations expose the exact exclusions; no silent traversal.
    english = public_result(result)
    norwegian = visning.md_import(result)
    for path in (unsupported, nested):
        assert any(r['path'] == str(path) for r in english['skipped'])
        assert str(path) in norwegian
    assert len(store.dokumenter(project['id'])) == 1
    again = tjeneste.importer_dokumenter(store, project['id'], [str(sources)])
    assert next(r for r in again['resultater'] if 'dokument' in r)['nytt'] is False
    assert {p: sha256_fil(p) for p in before} == before

    # A reported subfolder can be selected explicitly without reimporting its parent.
    selected = tjeneste.importer_dokumenter(store, project['id'], [str(nested)])
    assert not selected['skipped'] and selected['resultater'][0]['nytt']
    assert len(store.dokumenter(project['id'])) == 2


def test_folder_with_no_supported_files_still_reports_exclusions(tmp_path):
    store = Lager(tmp_path/'data'); project = store.opprett_prosjekt('Scope')
    sources = tmp_path/'sources'; sources.mkdir()
    (sources/'data.json').write_text('{}', encoding='utf-8')
    (sources/'appendices').mkdir()
    result = tjeneste.importer_dokumenter(store, project['id'], [str(sources)])
    assert len(result['resultater']) == 1 and 'feil' in result['resultater'][0]
    assert {r['reason'] for r in result['skipped']} == {'unsupported_format', 'subdirectory'}
    assert not store.dokumenter(project['id'])


def test_explicit_children_are_not_also_reported_as_skipped(tmp_path):
    store = Lager(tmp_path/'data'); project = store.opprett_prosjekt('Scope')
    sources = tmp_path/'sources'; sources.mkdir()
    (sources/'report.txt').write_text(QUOTE, encoding='utf-8')
    nested = sources/'earlier-years'; nested.mkdir()
    (nested/'older.txt').write_text('Older report.', encoding='utf-8')
    unsupported = sources/'data.json'; unsupported.write_text('{}', encoding='utf-8')
    result = tjeneste.importer_dokumenter(store, project['id'],
                                        [str(sources), str(nested), str(unsupported)])
    assert not result['skipped']
    assert len(store.dokumenter(project['id'])) == 2
    assert len([r for r in result['resultater'] if 'feil' in r]) == 1
    assert next(r['sti'] for r in result['resultater'] if 'feil' in r) == str(unsupported)


def test_identical_bytes_with_different_formats_keep_distinct_interpretations(tmp_path):
    store=Lager(tmp_path/'data'); project=store.opprett_prosjekt('Files')
    files = [tmp_path/'data.txt', tmp_path/'data.csv']
    for file in files: file.write_text('field,value\npolicy,yes\n', encoding='utf-8')
    docs = [importer_dokument(store,project['id'],file)[0] for file in files]
    assert docs[0]['sha256'] == docs[1]['sha256'] and docs[0]['id'] != docs[1]['id']
    assert docs[0]['source_metadata']['format'] == 'txt' and docs[1]['source_metadata']['format'] == 'csv'
    assert importer_dokument(store,project['id'],files[1])[0]['id'] == docs[1]['id']
