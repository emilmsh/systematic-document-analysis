"""Independent task execution, explicit input limits and long-path regression."""
import json
import os
from pathlib import Path
import zipfile

import pytest
from openpyxl import load_workbook

from kildeanalyse import tjeneste
from kildeanalyse.adaptere import ADAPTERE
from kildeanalyse.execution import prepare, preview, execute, settings
from kildeanalyse.file_io import copy_file, native_path
from kildeanalyse.modell import Plan, Motorsvar
from kildeanalyse.prompt import bygg_inputpakke
from test_tasks import fixture


def input_for(tmp_path, *, file_tools):
    source = tmp_path/'large.txt'
    source.write_text('Complete original input. ' * 10000, encoding='utf-8')
    document = {'id':'d1', 'navn':'large.txt', 'sha256':'hash', 'antall_sider':1,
                'lagret_kopi':str(source), 'sider':[{'nr':1, 'tekst':source.read_text(), 'tegn':source.stat().st_size}]}
    plan = Plan('Read', task_instructions='Count every occurrence.', motor='claude_cli',
                motorinnstillinger={'file_tools':file_tools, 'input_budget_bytes':8000})
    package = bygg_inputpakke(plan, document, forsok_id='f1', kjoring_id='k1')
    return plan, document, package


def test_large_cli_file_keeps_one_worker_and_complete_source(tmp_path):
    plan, document, full = input_for(tmp_path, file_tools=True)
    sent = []
    class Reader:
        def kjor(self, package, *args):
            sent.append(package)
            return Motorsvar('', {'result':10000, 'source_units_read':[1], 'limitations':[]})
    inspection = preview(plan, document, full)
    assert inspection['processing']['mode'] == 'file'
    execute(plan, document, full, Reader(), lambda:False, tmp_path)
    assert len(sent) == 1
    assert sent[0].sider == full.sider
    assert sent[0].local_source_path == full.local_source_path
    assert sent[0].systeminstruks == full.systeminstruks
    assert sent[0].svarskjema == full.svarskjema
    assert sent[0].hash() == inspection['input_hash']
    assert 'SOURCE_GUIDE.md' in sent[0].brukermelding


def test_oversized_text_input_fails_without_dispatch_or_summarisation(tmp_path):
    plan, document, package = input_for(tmp_path, file_tools=False)
    class Reader:
        def kjor(self, *args): pytest.fail('Must not dispatch')
    with pytest.raises(ValueError, match='No task execution'):
        execute(plan, document, package, Reader(), lambda:False, tmp_path)
    with pytest.raises(ValueError, match='removed'):
        settings({'document_processing':'auto'})


def test_each_iteration_has_fresh_adapter_and_no_previous_result(tmp_path, monkeypatch):
    store, aid, runs, _ = fixture(tmp_path, monkeypatch)
    workers = []
    def read(self, package, *args):
        assert not hasattr(self, 'last_answer')
        assert 'ANSWER-FROM-PREVIOUS-FILE' not in package.brukermelding
        assert 'ANSWER-FROM-PREVIOUS-FILE' not in package.systeminstruks
        self.last_answer = 'ANSWER-FROM-PREVIOUS-FILE'
        workers.append(self)
        reply = {'result':self.last_answer, 'source_units_read':[s.nr for s in package.sider], 'limitations':[]}
        return Motorsvar(json.dumps(reply), reply)
    monkeypatch.setattr(ADAPTERE['claude_cli'], 'kjor', read)
    tjeneste.godkjenn_plan(store, aid, 'Test fixture')
    report = tjeneste.start(store, aid)
    assert not report['run_issues']
    assert len(workers) == len(runs) == len({id(w) for w in workers})


@pytest.mark.skipif(os.name != 'nt', reason='Windows CopyFile2 regression')
def test_copy_and_complete_export_with_paths_over_260_characters(tmp_path, monkeypatch):
    deep = tmp_path/('a'*65)/('b'*65)
    deep.mkdir(parents=True)
    original = deep/('source'*12+'.txt')
    native_path(original).write_text('Unicode æøå and exact bytes.\n', encoding='utf-8')
    target = deep/('target'*12+'.txt')
    assert len(str(original)) > 260 and len(str(target)) > 260
    copy_file(original, target)
    assert native_path(target).read_bytes() == native_path(original).read_bytes()
    # Exercise the actual exporter, not just the low-level copy helper.
    visible = tmp_path / ('v' * max(1, 165 - len(str(tmp_path))))
    monkeypatch.setenv('SDA_PROJECTS_ROOT', str(visible))
    store, aid, runs, _ = fixture(tmp_path, monkeypatch)
    tjeneste.godkjenn_plan(store, aid, 'Test fixture')
    tjeneste.start(store, aid)
    output = tjeneste.eksporter(store, aid, output_directory=str(visible))
    assert load_workbook(native_path(output['workbook']))['Results'].max_row == 3
    with zipfile.ZipFile(native_path(output['documentation_archive'])) as archive:
        assert any(len(str(Path(output['mappe']) / n)) > 260 for n in archive.namelist())
        for run in runs:
            doc = store.dokument(run['dokument_id'])
            source = next(n for n in archive.namelist() if n.startswith('sources/'+doc['id']+'_'))
            assert archive.read(source) == Path(doc['lagret_kopi']).read_bytes()
            attempt = store.forsok_for_kjoring(run['id'])[0]
            assert archive.read(f'audit/attempts/{attempt["id"]}/raasvar.txt').decode('utf-8') == attempt['raasvar']


@pytest.mark.parametrize('engine', ['openai_api', 'anthropic_api', 'openrouter_api',
                                   'kompatibel_api', 'azure_foundry_api'])
def test_oversized_api_file_is_visible_in_status_and_workbook_and_next_file_runs(tmp_path, monkeypatch, engine):
    store, aid, initial, sent = fixture(tmp_path, monkeypatch, engine=engine, large=True)
    project_id = store.analyse(aid)['prosjekt_id']
    small = tmp_path/'small.txt'
    small.write_text('A small source with enough text to read.', encoding='utf-8')
    imported = tjeneste.importer_dokumenter(store, project_id, [str(small)])
    small_id = imported['resultater'][0]['dokument']['id']
    tjeneste.ny_planversjon(store, aid, 'Small input budget', motorinnstillinger={'input_budget_bytes':8000})
    runs = tjeneste.legg_til_kjoringer(store, aid, [initial[0]['dokument_id'], small_id])['nye']
    tjeneste.godkjenn_plan(store, aid, 'Test fixture')
    report = tjeneste.start(store, aid, kjoring_ider=[r['id'] for r in runs])
    failed, successful = runs
    assert report['workflow_block'] is None
    assert len(sent) == 1 and sent[0].dokument_id == small_id  # no dispatch for the oversized file
    assert store.kjoring(failed['id'])['status'] == 'feilet'
    assert store.kjoring(successful['id'])['status'] == 'fullført'
    assert not store.forsok_for_kjoring(failed['id'])
    status = tjeneste.vis_status(store, aid, details=False)
    issue, = status['run_issues']
    assert issue['run_id'] == failed['id'] and 'input_budget_bytes' in issue['reason']
    output = tjeneste.eksporter(store, aid)
    book = load_workbook(output['workbook'])
    def records(sheet):
        values = list(book[sheet].values)
        return [dict(zip(values[0], row)) for row in values[1:]]
    row = next(r for r in records('Results') if r['Run'] == failed['id'])
    assert all(value is None for key, value in row.items() if key not in ('Run', 'Document', 'Status', 'Human review'))
    assert 'feilet' in row['Status']
    assert any(r['Run'] == failed['id'] and 'input_budget_bytes' in r['Message']
               for r in records('Errors and notes'))
    book.close()
