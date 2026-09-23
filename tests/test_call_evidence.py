import copy
import json
from pathlib import Path
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

from kildeanalyse import tjeneste
from kildeanalyse.call_evidence import call_records, call_warnings
from kildeanalyse.english_tools import register
from kildeanalyse.lager import Lager
from kildeanalyse.adaptere.simulert import SimulertAdapter

FIX = Path(__file__).parent/'fixtures/syntetisk'
CACHE = 'ERROR failed to load models cache: missing field supports_parallel_tool_calls'


def test_single_call_preserves_evidence_without_inferred_telemetry():
    manifest = {'startet':'start', 'avsluttet':'end', 'sesjon_id':'session-1',
                'kjoreparametre':{'modell':'requested', 'tenkenivaa':'high'},
                'forbruk':{'usage':{'input_tokens':30, 'output_tokens':3}},
                'motorinfo':{'returkode':0, 'varighet_sek':1, 'stderr':CACHE}}
    attempt = {'id':'kj1.f1', 'kjoring_id':'kj1', 'motor':'codex_cli', 'simulert':False,
               'input_sti':'attempts/kj1.f1', 'manifest_json':json.dumps(manifest)}
    original = copy.deepcopy(attempt)
    records = call_records(attempt, 'source.pdf')
    row, = records
    assert row['input_tokens'] == 30 and row['session_id'] == 'session-1'
    assert row['reported_model'] is None and row['reported_effort'] is None
    assert row['requested_effort'] == 'high' and row['exit_code'] == 0
    assert row['started_at'] == 'start' and row['artifact_subdirectory'] == ''
    assert [w['code'] for w in call_warnings(records)] == ['CLI_MODEL_CACHE']
    assert attempt == original


@pytest.mark.parametrize('usage', [
    {'input_tokens':0, 'output_tokens':5, 'cache_read_input_tokens':0},
    {'prompt_tokens':0, 'completion_tokens':5, 'prompt_tokens_details':{'cached_tokens':0}},
])
def test_missing_and_zero_usage_are_distinct_and_pre_dispatch_failure_is_not_success(usage):
    stopped = {'dispatched':False, 'avsluttet':'end', 'feil':'Stopped before dispatch', 'motorinfo':{}}
    row, = call_records({'id':'f1', 'manifest_json':json.dumps(stopped)})
    assert row['status'] == 'not_dispatched' and row['input_tokens'] is None and row['exit_code'] is None
    success = {'forbruk':{'usage':usage}, 'modell_rapportert':'actual-model',
               'motorinfo':{'http_status':200, 'stderr':'diagnostic '+('x'*1000)}}
    row, = call_records({'id':'f2', 'manifest_json':json.dumps(success)})
    assert row['input_tokens'] == 0 and row['cached_input_tokens'] == 0
    assert row['output_tokens'] == 5 and row['reported_model'] == 'actual-model'
    assert row['warnings'][0]['truncated'] and len(row['warnings'][0]['message']) == 800
    assert call_records({'id':'not-started'}) == []


@pytest.mark.parametrize('language', ['nb','en'])
def test_warnings_do_not_block_queue_and_call_export_preserves_all_attempts(tmp_path, monkeypatch, language):
    store = Lager(tmp_path/'data')
    project = tjeneste.opprett_prosjekt(store,'Evidence',str(tmp_path/'visible'))
    tjeneste.importer_dokumenter(store, project['id'], [str(FIX/'fjordblikk_2025.pdf'),str(FIX/'steinbukk_2025.pdf')])
    result = tjeneste.opprett_analyse(store,project['id'],'Evidence','Read',
                                     motor='simulert',sprak=language)
    aid = result['analyse']['id']
    runs = tjeneste.legg_til_kjoringer(store,aid)['nye']
    tjeneste.godkjenn_plan(store,aid,'Test person')
    original = SimulertAdapter.kjor
    def read(self, *args, **kwargs):
        reply = original(self, *args, **kwargs)
        reply.motorinfo = {'returkode':0,'stderr':CACHE,'cli_versjon':'test-double',
                           'tenkenivaa_onsket':'high','tenkenivaa_rapportert':'ukjent'}
        reply.modell_rapportert = None
        reply.forbruk = {'usage':{'input_tokens':100,'output_tokens':10}}
        return reply
    monkeypatch.setattr(SimulertAdapter,'kjor',read)
    tjeneste.start(store,aid)
    status = tjeneste.vis_status(store,aid,details=False)
    assert status['teller'] == {'fullført':2}
    assert status['workflow_block'] is None and status['run_issues'] == []
    assert len(status['run_warnings']) == 2
    assert all('sider' not in row['dokument'] for row in status['rader'])
    assert all('manifest_json' not in row['siste_forsok'] for row in status['rader'])
    functions = {}
    class Server:
        def tool(self, **kwargs):
            def register_fn(fn):
                functions[fn.__name__] = fn
                return fn
            return register_fn
    register(Server(),lambda:store)
    compact = json.loads(functions['show_status'](aid))
    full = json.loads(functions['show_status'](aid,details=True))
    assert compact['detail_level'] == 'compact' and full['detail_level'] == 'full'
    assert compact['teller'] == full['teller']
    assert len(json.dumps(compact)) < len(json.dumps(full))/2
    assert 'sider' in full['rader'][0]['document']
    # A new attempt keeps historical evidence and exposes each warning exactly once.
    tjeneste.nytt_forsok(store,runs[0]['id'],'Explicit retry in local test')
    tjeneste.start(store,aid)
    export = tjeneste.eksporter(store,aid)
    root = Path(export['mappe'])
    book = load_workbook(export['workbook'])
    import zipfile
    assert book['Resultater' if language == 'nb' else 'Results'].max_row == 3
    with zipfile.ZipFile(export['documentation_archive']) as archive:
        records = []
        for run in runs:
            audit = json.loads(archive.read(f'audit/{run["id"]}.json'))
            for detail in audit['forsok']:
                records.extend(detail['model_calls'])
            for attempt in store.forsok_for_kjoring(run['id']):
                original = Path(attempt['input_sti'])/'manifest.json'
                assert archive.read(f'audit/attempts/{attempt["id"]}/manifest.json') == original.read_bytes()
        assert len(records) == 3
        assert all(CACHE in json.dumps(r) for r in records)
