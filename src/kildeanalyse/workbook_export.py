"""Portable XLSX export for the Python plugin, using its existing openpyxl runtime."""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .call_evidence import reported


def write_workbook(directory: Path, analysis: dict, versions: list[dict], runs: list[dict],
                   attempts: list[dict], reviews: list[dict], language: str,
                   include_sources: bool, document_folder: str, *, calls: list[dict] | None = None) -> tuple[str, list[str]]:
    nb = language == 'nb'
    choose = lambda no, en: no if nb else en
    book = Workbook()
    book.remove(book.active)
    notices = []
    overview_name = choose('Oversikt', 'Overview')
    documentation = directory / document_folder
    calls = calls or []
    not_reported = choose('Ikke rapportert', 'Not reported')

    def sheet(name, headers, rows, widths=None):
        if len(rows) > 1048575:
            raise ValueError(f'{name}: too many rows for one Excel sheet; use the legacy CSV export.')
        ws = book.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([''] * len(headers))
            for col, value in enumerate(row, 1):
                cell = ws.cell(ws.max_row, col)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                if value is None:
                    value = ''
                if isinstance(value, str):
                    if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
                        digest = hashlib.sha256(value.encode('utf-8')).hexdigest()
                        text_path = documentation/'text'/f'{digest[:24]}.txt'
                        text_path.parent.mkdir(parents=True, exist_ok=True)
                        if text_path.exists() and text_path.read_text(encoding='utf-8') != value:
                            raise ValueError('Text attachment hash collision; export was not published.')
                        text_path.write_text(value, encoding='utf-8')
                        cell.hyperlink = quote(text_path.relative_to(directory).as_posix())
                        notices.append(f'{name}!{cell.coordinate}')
                        value = choose('Teksten vises i vedlagt tekstfil (Excel-begrensning).',
                                       'Full text is in the linked text file (Excel limitation).')
                    cell.value = value
                    cell.data_type = 's'  # Source strings such as =HYPERLINK(...) are never formulas.
                else:
                    cell.value = value
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='0F5F91')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[1].height = 32
        for index in range(1, len(headers)+1):
            ws.column_dimensions[get_column_letter(index)].width = (widths or {}).get(index, 24)
        for row in ws.iter_rows(min_row=2):
            # Set a readable initial height; Excel can expand unusually long rows further.
            lines = max(sum(max(1, math.ceil(len(part) / max(8, ws.column_dimensions[cell.column_letter].width - 2)))
                            for part in str(cell.value or '').split('\n')) for cell in row)
            ws.row_dimensions[row[0].row].height = min(300, max(30, lines * 15 + 8))
            for cell in row:
                if cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='EEF4F8')
                if cell.value in ('feilet', 'valideringsfeil', 'uavklart', 'avvist'):
                    cell.font = Font(color='A61B1B', bold=True)
        return ws

    plan = versions[-1]['plan']
    summary = [
        [choose('Analyse', 'Analysis'), analysis['navn']],
        [choose('Formål', 'Purpose'), plan.formaal],
        [choose('Planversjon', 'Plan version'), versions[-1]['versjon']],
        [choose('Planstatus', 'Plan status'), versions[-1]['status']],
        [choose('Plan godkjent av', 'Plan approved by'), versions[-1].get('godkjent_av') or '—'],
        [choose('Lesemotor', 'Reader'), plan.motor],
        [choose('Bestilt modell', 'Requested model'), plan.modell],
        [choose('Bestilt tenkenivå', 'Requested reasoning effort'), plan.motorinnstillinger.get('tenkenivaa', '')],
        [choose('Språk', 'Language'), plan.sprak],
        [choose('Dokumentkjøringer', 'Document runs'), len(runs)],
        [choose('Registrerte lesekall', 'Recorded reader calls'), len(calls)],
        [choose('Tekniske advarsler', 'Technical warnings'), sum(len(c['warnings']) for c in calls)],
        [choose('Kontroll', 'Review'), choose('Automatisk validering er ikke menneskelig kontroll. Excel-endringer føres ikke tilbake.',
                                             'Automatic validation is not human review. Excel edits do not write back.')],
        [choose('Dokumentasjon', 'Documentation'), f'{document_folder}/analyse.json'],
        [choose('Modellinnstillinger', 'Model settings'), choose('Bestilte innstillinger er ikke bevis for faktisk rapportert modell/tenkenivå. Se Kjøringer og modellkall.',
                                                                'Requested settings do not prove the reported model/effort. See Runs and model calls.')],
    ]
    kinds = {bool(r['gjeldende_forsok']['simulert']) for r in runs if r.get('gjeldende_forsok')}
    summary.append([choose('Kjøringstype', 'Run type'),
                    choose('SIMULERT', 'SIMULATED') if kinds == {True} else
                    choose('Blandet: simulert og ekte', 'Mixed: simulated and real') if len(kinds)>1 else
                    choose('Ekte modellkall', 'Real model calls') if kinds else choose('Ikke startet', 'Not started')])
    results, evidence = [], []
    for run in runs:
        doc, kj = run['dokument'], run['kjoring']
        attempt = run.get('gjeldende_forsok') or {}
        validation = json.loads(attempt.get('validering_json') or '{}')
        warnings = [w.get('melding', '') for w in validation.get('advarsler', [])]
        source = run.get('source_metadata', {})
        warnings += source.get('ocr', {}).get('warnings', [])
        warnings += [f"{w['code']} (call {c['call_index']}): {w['message']}"
                     for c in calls if c['attempt_id'] == attempt.get('id') for w in c['warnings']]
        assessment = run.get('vurderinger') or {}
        if not assessment:
            results.append([doc['navn'], '', '', '', kj['status'], '', '', kj['planversjon_id'], kj['id'],
                            '\n'.join(filter(None, [attempt.get('feil') or kj.get('merknad'), *warnings]))])
        for kid, value in assessment.items():
            errors = '; '.join(str(e) for e in value.get('valideringsfeil', []))
            results.append([doc['navn'], kid, value['svar'], value.get('kommentar', ''), kj['status'],
                            'ok' if value.get('validering_gyldig') else errors or choose('Feil', 'Invalid'),
                            value['kontrollstatus'], kj['planversjon_id'], kj['id'], '\n'.join(warnings)])
            for item in value['belegg']:
                evidence.append([doc['navn'], kid, value['svar'], item.get('sitat', ''),
                                 item.get('source', {}).get('location', str(item.get('side', ''))),
                                 item.get('side'), kj['id'], attempt.get('id', ''), value.get('kilde', '')])
    sheet(overview_name, [choose('Felt', 'Field'), choose('Verdi', 'Value')], summary, {1:30, 2:100})
    sheet(choose('Resultater', 'Results'),
          ['Dokument','Kriterium','Svar','Begrunnelse','Kjøringsstatus','Validering','Menneskelig kontroll','Planversjon','Kjøring','Merknader'] if nb else
          ['Document','Criterion','Answer','Comment','Run status','Validation','Human review','Plan version','Run','Notes'],
          results, {1:32, 4:80, 10:80})
    ev = sheet(choose('Belegg', 'Evidence'),
               ['Dokument','Kriterium','Svar','Sitat','Kildeplassering','Kildeenhet','Kjøring','Forsøk','Svargrunnlag'] if nb else
               ['Document','Criterion','Answer','Quote','Source location','Source unit','Run','Attempt','Answer source'],
               evidence, {1:32, 4:100, 5:35})
    source_folder = choose('Kilder', 'Sources')
    if include_sources:
        # Use run IDs, not filenames, to disambiguate sources with the same name.
        docs_by_run = {r['kjoring']['id']:r['dokument'] for r in runs}
        for row in range(2, ev.max_row+1):
            doc = docs_by_run[ev.cell(row,7).value]
            filename = f'{doc["id"]}_{doc["navn"]}'
            if (directory/source_folder/filename).is_file():
                ev.cell(row,1).hyperlink = quote(f'{source_folder}/{filename}')
                ev.cell(row,1).font = Font(color='0563C1', underline='single')
    review_keys = ['tid','ansvarlig','handling','kriterium_id','begrunnelse','forsok_id','opprinnelig','nytt']
    sheet(choose('Kontroll','Review'),
          ['Tid','Ansvarlig','Handling','Kriterium','Begrunnelse','Forsøk','Opprinnelig','Nytt'] if nb else
          ['Time','Reviewer','Action','Criterion','Reason','Attempt','Original','New'],
          [[r.get(k, '') for k in review_keys] for r in reviews], {5:75,7:65,8:65})
    attempt_keys = ['id','kjoring_id','status','startet','avsluttet','motor','simulert','modell_onsket','modell_rapportert','tenkenivaa_onsket','feil']
    sheet(choose('Kjøringer','Runs'),
          ['Forsøk','Kjøring','Status','Startet','Avsluttet','Lesemotor','Simulert','Bestilt modell','Rapportert modell','Bestilt tenkenivå','Feil'] if nb else
          ['Attempt','Run','Status','Started','Finished','Reader','Simulated','Requested model','Reported model','Requested effort','Error'],
          [[(reported(r.get(k)) or not_reported) if k == 'modell_rapportert' else r.get(k, '')
            for k in attempt_keys] for r in attempts], {11:80})
    call_keys = ['document_name','attempt_id','call_index','stage','status','reader','simulated',
                 'session_id','cli_version','requested_model','reported_model','requested_effort','reported_effort',
                 'started_at','finished_at','duration_seconds','exit_code','input_tokens','output_tokens',
                 'cached_input_tokens','process_id','request_id','http_status']
    call_rows = []
    for c in calls:
        row = [not_reported if c.get(k) is None else c[k] for k in call_keys]
        row += ['\n'.join(f"{w['code']}: {w['message']}" for w in c['warnings']), c.get('error') or '']
        row += [''] * 3  # portable artifact links are added below, only when files exist
        call_rows.append(row)
    call_sheet = sheet(choose('Modellkall','Model calls'),
          (['Dokument','Forsøk','Kall','Trinn','Status','Lesemotor','Simulert','Sesjons-ID','CLI-versjon',
            'Bestilt modell','Rapportert modell','Bestilt tenkenivå','Rapportert tenkenivå','Startet','Avsluttet',
            'Varighet (sek.)','Returkode','Inputtokens','Outputtokens','Cachetokens','Prosess-ID','Request-ID',
            'HTTP-status','Advarsler','Feil','Input','Råsvar','Manifest'] if nb else
           ['Document','Attempt','Call','Stage','Status','Reader','Simulated','Session ID','CLI version',
            'Requested model','Reported model','Requested effort','Reported effort','Started','Finished',
            'Duration (sec.)','Exit code','Input tokens','Output tokens','Cached tokens','Process ID','Request ID',
            'HTTP status','Warnings','Error','Input','Raw reply','Manifest']),
          call_rows, {1:32, 3:10, 8:42, 14:36, 15:36, 24:80, 25:65})
    for row, call in enumerate(calls, 2):
        artifact = Path(document_folder)/'modellkall'/call['attempt_id']/call['artifact_subdirectory']
        for col, name in ((26,'input.json'), (27,'raasvar.txt'), (28,'manifest.json')):
            path = artifact/name
            cell = call_sheet.cell(row,col)
            if (directory/path).is_file():
                cell.value = name
                cell.hyperlink = quote(path.as_posix())
                cell.font = Font(color='0563C1', underline='single')
            else:
                cell.value = choose('Ikke tilgjengelig','Not available')
    if notices:
        book[overview_name].append([choose('Tekstvedlegg','Text attachments'), ', '.join(notices)])
    filename = choose('Resultater.xlsx', 'Results.xlsx')
    book.save(directory/filename)
    book.close()
    return filename, notices
