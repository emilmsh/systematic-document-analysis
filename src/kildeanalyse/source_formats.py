"""File adapters produce numbered source units, never fictitious PDF pages.

The legacy response fields side/sider_lest carry unit IDs. Locators are resolved
from the preserved source, not invented by the reading model.
"""
import csv
import io

SUPPORTED = {'.pdf', '.docx', '.xlsx', '.csv', '.tsv', '.txt', '.md'}


def unit(number, text, kind, location, **coordinates):
    return {'nr':number, 'tekst':text, 'tegn':len(''.join(text.split())),
            'source':dict(kind=kind, location=location, **coordinates), 'readable':bool(text.strip())}


def location(document, number):
    item = next((s for s in document['sider'] if s['nr'] == number), None)
    if item is None:
        return {'kind':'unknown', 'location':f'Unknown unit {number}'}
    return item.get('source') or {'kind':'pdf_page', 'location':f'PDF page {number}', 'page':number}


def decode_text(path):
    data = path.read_bytes()
    # Explicit BOM handling; never silently replace undecodable bytes.
    return data.decode('utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')


def extract(path):
    suffix = path.suffix.lower()
    if suffix in ('.txt', '.md'):
        lines = decode_text(path).splitlines()
        units = [unit(i, line, 'line', f'Line {i}', line=i) for i, line in enumerate(lines, 1) if line.strip()]
        # IDs are consecutive; physical line positions stay in source metadata.
        for i, row in enumerate(units, 1): row['nr'] = i
        return units, {'format':suffix[1:], 'scope':'Non-empty text lines; Markdown is read as source text.', 'structure':{}}, 'stdlib text UTF-8/UTF-16'
    if suffix in ('.csv', '.tsv'):
        text = decode_text(path)
        if suffix == '.tsv':
            dialect = 'excel-tab'
        elif not any(delimiter in text for delimiter in ',;\t|'):
            dialect = 'excel'
        else:
            try:
                dialect = csv.Sniffer().sniff(text[:65536], delimiters=',;\t|')
            except csv.Error:
                # Sniffer can fail on valid quoted multiline fields. Parse candidates;
                # accept only an unambiguous rectangular table, never discard records.
                candidates = []
                for delimiter in ',;\t|':
                    try:
                        sample = list(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, strict=True))
                    except csv.Error:
                        continue
                    widths = {len(row) for row in sample if row}
                    if len(widths) == 1 and next(iter(widths)) > 1: candidates.append(delimiter)
                if len(candidates) != 1:
                    raise ValueError('Ambiguous CSV delimiter; save as UTF-8 TSV or consistently delimited CSV.')
                dialect = type('DetectedDialect', (csv.excel,), {'delimiter':candidates[0]})
        rows = csv.reader(io.StringIO(text, newline=''), dialect=dialect, strict=True)
        units, headers = [], []
        for index, row in enumerate(rows, 1):
            if index == 1: headers = row
            if any(value.strip() for value in row):
                rendered = '\n'.join(f'Column {c}: {value}' for c, value in enumerate(row, 1))
                units.append(unit(len(units)+1, rendered, 'record', f'Record {index}', record=index, cells=row))
        return units, {'format':suffix[1:], 'scope':'All non-empty records including the first record; values remain text.',
                       'structure':{'first_record':headers}}, 'stdlib csv'
    if suffix == '.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        import docx
        document = Document(path)
        units, headings = [], []
        def walk(container, prefix):
            for index, item in enumerate(container.iter_inner_content(), 1):
                ref = f'{prefix}/block {index}'
                if isinstance(item, Paragraph):
                    if item.text.strip():
                        units.append(unit(len(units)+1, item.text, 'paragraph', ref, block=ref))
                        if item.style and item.style.name.startswith('Heading'): headings.append(item.text)
                else:
                    seen = set()
                    for r, row in enumerate(item.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            if cell._tc in seen: continue  # merged cells share an XML element
                            seen.add(cell._tc)
                            walk(cell, f'{ref}/table row {r} cell {c}')
        walk(document, 'Body')
        notes = ['Main document body and nested tables only. Headers, footers, footnotes, comments, tracked-change wrappers, text boxes and images are not analysed.',
                 'Word pagination is layout-dependent; references identify blocks, not page numbers.']
        return units, {'format':'docx', 'scope':notes, 'structure':{'headings':headings}}, f'python-docx {docx.__version__}'
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        import openpyxl
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        values = None
        try:
            values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            units, sheets, missing = [], [], []
            for sheet in formulas:
                cached = values[sheet.title]
                sheet.reset_dimensions(); cached.reset_dimensions()
                first_row = None
                for row, cached_row in zip(sheet.iter_rows(), cached.iter_rows()):
                    cells, cell_text = [], []
                    for cell, value_cell in zip(row, cached_row):
                        if cell.value is None: continue
                        value = cell.value
                        entry = {'cell':cell.coordinate, 'value':str(value), 'number_format':cell.number_format}
                        rendered = f'{cell.coordinate}: {value}'
                        if cell.data_type == 'f':
                            entry = {'cell':cell.coordinate, 'formula':str(value), 'cached_value':None if value_cell.value is None else str(value_cell.value), 'number_format':cell.number_format}
                            rendered = f'{cell.coordinate}: formula {value}; cached value: {value_cell.value if value_cell.value is not None else "UNAVAILABLE"}'
                            if value_cell.value is None: missing.append(f'{sheet.title}!{cell.coordinate}')
                        cells.append(entry); cell_text.append(rendered)
                    if not cells: continue
                    if first_row is None: first_row = cells
                    row_number = next(c.row for c in row if c.value is not None)
                    span = f"{cells[0]['cell']}:{cells[-1]['cell']}"
                    units.append(unit(len(units)+1, '\n'.join(cell_text), 'sheet_row',
                                      f'{sheet.title}!{span}', sheet=sheet.title, row=row_number, cells=cells,
                                      sheet_state=sheet.sheet_state))
                sheets.append({'sheet':sheet.title,'state':sheet.sheet_state,'first_nonempty_row':first_row})
            return units, {'format':'xlsx', 'scope':['All non-empty rows on all sheets, including hidden sheets/rows.',
                'Formulas are not recalculated. Cached values may be missing or stale. Formatting is not a rendered display.',
                'Charts, images, comments and other embedded objects are not analysed.'],
                'missing_formula_values':missing, 'structure':{'sheets':sheets}}, f'openpyxl {openpyxl.__version__}'
        finally:
            formulas.close()
            if values is not None: values.close()
    raise ValueError(f'Unsupported format: {suffix}')


def metadata(document):
    return document.get('source_metadata') or {'format':'pdf', 'scope':'Extracted PDF text; physical pages, no OCR.', 'structure':{}}


def annotate_assessments(document, assessments):
    if assessments:
        for assessment in assessments.values():
            for evidence in assessment.get('belegg', []):
                evidence['source_unit'] = evidence.get('side')
                evidence['source'] = location(document, evidence.get('side'))
    return assessments
