"""Human-readable datasets for arbitrary task results, using the plugin's XLSX runtime."""
from __future__ import annotations

import csv
import json
import math
import re
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .task_dataset import add_result, describe, pointer, source_locations


def write_workbook(directory, versions, records, *, include_csv=False,
                   csv_directory=None, list_layout='sheets', row_scope='runs'):
    if list_layout not in ('sheets', 'inline'):
        raise ValueError('list_layout must be sheets or inline.')
    nb = versions[-1]['plan'].sprak == 'nb'
    tr = lambda no, en: no if nb else en
    row_unit = tr('Én kjøring', 'One run') if row_scope == 'runs' else tr('Én dokumentrad', 'One document row')
    book = Workbook()
    book.remove(book.active)
    used_names, overflow = set(), []
    dictionary, datasets, issues = [], OrderedDict(), []

    def name_for(name):
        base = re.sub(r'[\\/*?:\[\]\x00-\x1f]', '_', name).strip(" '")[:31] or 'Data'
        candidate, index = base, 2
        while candidate.casefold() in used_names:
            suffix = f' {index}'
            candidate = base[:31-len(suffix)] + suffix
            index += 1
        used_names.add(candidate.casefold())
        return candidate

    dictionary_name = name_for(tr('Variabler', 'Variables'))
    texts_name = name_for(tr('Lange tekster', 'Long texts'))
    results_name = name_for(tr('Resultater', 'Results'))

    def write_cell(cell, value):
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        if isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15:
            value = str(value)  # Excel's numeric precision must not change identifiers or exact integers.
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError('Non-finite numbers cannot be exported as dataset values.')
        if isinstance(value, str):
            if ILLEGAL_CHARACTERS_RE.search(value):
                value = ILLEGAL_CHARACTERS_RE.sub(lambda m: f'\\u{ord(m[0]):04x}', value)
            # Keep long text inside the same workbook, with no truncation or loose attachments.
            if len(value) > 30000 or (len(value) > 2200 and cell.parent.title != texts_name):
                parts = [value[i:i+2000] for i in range(0, len(value), 2000)]
                first = len(overflow) + 2
                ref = f'{cell.parent.title}!{cell.coordinate}'
                overflow.extend([[ref, n, text] for n, text in enumerate(parts, 1)])
                value = tr('Hele teksten: ', 'Full text: ') + f'{texts_name} ({len(parts)} ' + tr('deler)', 'parts)')
                cell.hyperlink = f"#'{texts_name.replace(chr(39), chr(39)*2)}'!C{first}"
            cell.value = value
            cell.data_type = 's'  # Source text must never become an Excel formula.
        else:
            cell.value = value
            if isinstance(value, (int, float)):
                cell.number_format = '0' if isinstance(value, int) else '0.##########'

    def sheet(name, headers, rows, widths=None):
        if len(rows) > 1048575 or len(headers) > 16384:
            raise ValueError(f'{name}: dataset exceeds Excel limits; narrow the output contract or selected files.')
        ws = book.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = 'B2' if len(headers) > 5 else 'A2'
        for r, values in enumerate([headers, *rows], 1):
            for c, value in enumerate(values, 1):
                cell = ws.cell(r, c)
                write_cell(cell, value)
                cell.font = Font(name='Arial', size=11, color='243444')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if r == 1:
                    cell.fill = PatternFill('solid', fgColor='233F5B')
                    cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                elif r % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F0F4F7')
        for c in range(1, len(headers)+1):
            ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 24)
        header_lines = max((math.ceil(len(str(c.value or '')) / max(8, ws.column_dimensions[c.column_letter].width-2)) for c in ws[1]), default=1)
        ws.row_dimensions[1].height = max(30, min(100, header_lines*15+6))
        for row in ws.iter_rows(min_row=2):
            lines = max((sum(max(1, math.ceil(len(p) / max(8, ws.column_dimensions[c.column_letter].width-2)))
                             for p in str(c.value if c.value is not None else '').split('\n')) for c in row), default=1)
            ws.row_dimensions[row[0].row].height = min(409, max(22, lines*15+4))
        ws.auto_filter.ref = ws.dimensions
        return ws

    schema_groups = {}
    for record in records:
        run, doc, plan, chosen, view = (record[k] for k in ('run', 'document', 'plan', 'chosen', 'view'))
        aid = chosen['id'] if chosen else ''
        result = view['result'] if view else None
        validation = (view or {}).get('result_validation') or {}
        review = (view or {}).get('review_status', 'ikke kontrollert')
        accepted = bool(view and validation.get('gyldig') and review != 'avvist' and
                        (chosen['status'] == 'fullført' or view.get('result_origin') == 'human_correction'))
        if not view:
            state = tr('Ingen resultatleveranse', 'No result')
        elif review == 'avvist':
            state = tr('Avvist; ikke med i datasettet', 'Rejected; excluded from dataset')
        elif not accepted:
            state = tr('Ugyldig/uavklart; ikke med i datasettet', 'Invalid/uncertain; excluded from dataset')
        elif result == [] or result == {}:
            state = tr('Tomt resultat', 'Empty result')
        elif result is None:
            state = tr('Resultatet er null', 'Result is null')
        else:
            state = tr('Med i datasettet', 'Included in dataset')
        pid = run['planversjon_id']
        key = json.dumps(plan.output_schema, sort_keys=True, ensure_ascii=False)
        group = schema_groups.setdefault(key, pid)
        record['dataset_key'] = group
        tables = datasets.setdefault(group, describe(plan.output_schema))
        context = {'document': doc['navn'], 'document_id': doc['id'], 'run_id': run['id'], 'attempt_id': aid,
                   'plan_id': pid, 'source_sha256': doc['sha256'], 'review': review,
                   'simulated': bool(chosen and chosen['simulert']), 'origin': (view or {}).get('result_origin', ''),
                   'source': doc, 'plan': plan}
        if accepted:
            add_result(tables, result, context)
        record['dataset_included'] = accepted
        record['dataset_state'] = state
        errors = chosen.get('feil') if chosen else run.get('merknad')
        if not errors:
            errors = '\n'.join(e.get('melding', str(e)) for e in validation.get('feil', []))
        limitations = ((view or {}).get('response') or {}).get('limitations', [])
        # The validator also records limitations as a Python-list string. Their
        # readable source values already have a dedicated column and issue row.
        warnings = '\n'.join(x.get('melding', '') for x in validation.get('advarsler', [])
                             if x.get('type') != 'validation_scope' and
                             not (x.get('type') == 'limitations' and limitations))
        for kind, message in [(tr('Feil', 'Error'), errors),
                              (tr('Begrensning', 'Limitation'), '\n'.join(limitations)),
                              (tr('Validering', 'Validation'), warnings)]:
            if message:
                issues.append([doc['navn'], run['id'], aid, kind, message])

    def headers_for(table):
        names = set()
        headers = []
        for column in table.columns.values():
            label = column.schema.get('title') or '.'.join(column.path) or tr('Resultat', 'Result')
            if column.role == 'count':
                label += tr(' (antall)', ' (count)')
            base, i = str(label), 2
            while label.casefold() in names:
                label = f'{base} ({i})'
                i += 1
            names.add(label.casefold())
            headers.append(label)
        return headers

    meta = [('source_location', tr('Kildeplassering', 'Source location')),
            ('record_id', tr('Rad-ID', 'Record ID')), ('parent_id', tr('Overordnet rad-ID', 'Parent record ID')),
            ('review', tr('Menneskelig kontroll', 'Human review'))]

    # The primary dataset has exactly one row per selected run, including failed/unstarted runs.
    # Repeated values default to detail sheets; inline lists are an explicit option.
    primary_columns = []
    primary_headers = [tr('Dokument', 'Document'), tr('Status', 'Status'), tr('Datasett', 'Dataset'),
                       tr('Validering', 'Validation'), tr('Menneskelig kontroll', 'Human review'), tr('Kjøring', 'Run')]
    for pid, tables in datasets.items():
        for table in tables.values():
            if table.repeated and list_layout == 'sheets':
                if not table.path:
                    header = (f'{pid}: ' if len(datasets) > 1 else '') + tr('Resultatposter (antall)', 'Result items (count)')
                    primary_headers.append(header)
                    primary_columns.append((pid, table, None))
                    dictionary.append([results_name, header, '/', tr('Antall poster i detaljfanen.', 'Number of detail records.'),
                                       'integer', pid, '/', row_unit])
                continue
            for label, col in zip(headers_for(table), table.columns.values()):
                prefix = '.'.join(p for p in table.path if p != '*')
                header = f'{prefix}.{label}' if prefix else label
                if len(datasets) > 1:
                    header = f'{pid}: {header}'
                while header.casefold() in {h.casefold() for h in primary_headers}:
                    header += tr(' (variabel)', ' (variable)')
                primary_headers.append(header)
                primary_columns.append((pid, table, col))
                kind = col.schema.get('type', tr('Ikke deklarert', 'Not declared'))
                kind = ', '.join(kind) if isinstance(kind, list) else kind
                if col.role == 'count':
                    kind = 'integer'
                if table.repeated:
                    kind = tr('Liste av ', 'List of ') + str(kind)
                dictionary.append([results_name, header, pointer(col.path) or '/',
                    col.schema.get('description') or (tr('Antall elementer i listen.', 'Number of list items.') if col.role == 'count' else tr('Ikke beskrevet i planen.', 'Not described in the plan.')),
                    kind, pid, pointer(table.path) or '/', row_unit])
    primary_rows = []
    for record in records:
        run, chosen, view = record['run'], record['chosen'], record['view']
        values = []
        for pid, table, col in primary_columns:
            selected = [r for r in table.rows if r['run_id'] == run['id']]
            if pid != record['dataset_key'] or not record['dataset_included']:
                values.append(None)
                continue
            if col is None:
                values.append(None if view['result'] is None else len(selected))
                continue
            present = [r['values'].get(col.key) for r in selected]
            if table.repeated:
                # Number list items so multi-line passages cannot be confused with separate findings.
                values.append('\n\n'.join(f'{i}. ' + (tr('(mangler)', '(missing)') if col.key not in r['present'] else
                    '(null)' if v is None else str(v).lower() if isinstance(v, bool) else str(v))
                    for i, (r, v) in enumerate(zip(selected, present), 1)) if selected else None)
            else:
                values.append(present[0] if present else None)
        validation = (view or {}).get('result_validation') or {}
        validation_state = tr('Bestått', 'Passed') if validation.get('gyldig') else tr('Ikke bestått', 'Not passed') if validation else tr('Ikke kontrollert', 'Not checked')
        prefix = [record['document']['navn'], run['status'], record['dataset_state'], validation_state,
                  (view or {}).get('review_status', 'ikke kontrollert'), run['id']]
        primary_rows.append([*prefix, *values])
    widths = {1:32, 2:24, 3:38, 4:20, 5:22, 6:16}
    for i, (_, table, col) in enumerate(primary_columns, len(primary_headers)-len(primary_columns)+1):
        if col is None:
            widths[i] = 22
            continue
        sizes = [len(str(row[i-1] or '')) for row in primary_rows]
        widths[i] = min(90, max(22, max(sizes, default=0)*0.65)) if table.repeated or col.schema.get('type') not in ('integer','number','boolean') else 16
    primary = sheet(results_name, primary_headers, primary_rows, widths)
    book.move_sheet(primary, offset=-book.index(primary))
    if include_csv and csv_directory:
        csv_directory.mkdir(parents=True, exist_ok=True)
        with (csv_directory / f'{results_name}.csv').open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(primary_headers)
            writer.writerows(primary_rows)
    summaries = []
    for pid, tables in datasets.items():
        for table in tables.values():
            if not table.repeated or not table.rows:
                continue  # Already represented in the one-row-per-run primary dataset.
            label = table.schema.get('title') or (tr('Funn', 'Findings') if not table.path else '.'.join(table.path).replace('.*', ''))
            name = name_for(str(label))
            headers = headers_for(table)
            # Reserved metadata columns are visibly distinct even when a user variable has the same name.
            all_headers = [tr('Kildedokument', 'Source document')] + headers
            for _, title in meta:
                candidate = title
                while candidate.casefold() in {h.casefold() for h in all_headers}:
                    candidate += ' [SDA]'
                all_headers.append(candidate)
            if all_headers[0].casefold() in {h.casefold() for h in headers}:
                all_headers[0] += ' [SDA]'
            rows = []
            for row in table.rows:
                row['source_location'] = source_locations(table, row, row['plan'], row['source'])
                row['null_fields'] = '\n'.join(row['nulls'])
                row['missing_fields'] = '\n'.join(pointer(c.path) or '/' for c in table.columns.values() if c.key not in row['present'])
                rows.append([row['document']] + [row['values'].get(k) for k in table.columns] + [row[k] for k, _ in meta])
            widths = {1:30}
            for i, column in enumerate(table.columns.values(), 2):
                kind = column.schema.get('type')
                lengths = [len(str(r['values'].get(column.key, ''))) for r in table.rows]
                widths[i] = 14 if column.role == 'count' or kind in ('integer', 'number', 'boolean') else min(90, max(24, max(lengths, default=0)*0.7))
            ws = sheet(name, all_headers, rows, widths)
            summaries.append((name, pid, len(rows), table.schema.get('description') or
                              tr('Én resultatpost per kjøring' if table.repeated else 'Ett resultat per kjøring',
                                 'One result item per run' if table.repeated else 'One result per run')))
            for label, c in zip(headers, table.columns.values()):
                spec = c.schema
                kind = 'integer' if c.role == 'count' else spec.get('type', tr('Ikke deklarert', 'Not declared'))
                definition = tr('Antall elementer i listen. Ingen kombinasjon med andre lister.', 'Number of list items. No cross-product with other lists.') if c.role == 'count' else spec.get('description', tr('Ikke beskrevet i planen.', 'Not described in the plan.'))
                dictionary.append([name, label, pointer(c.path) or '/', definition, ', '.join(kind) if isinstance(kind, list) else kind,
                                   pid, pointer(table.path) or '/', summaries[-1][3]])
            if include_csv and csv_directory:
                csv_directory.mkdir(parents=True, exist_ok=True)
                with (csv_directory / f'{name}.csv').open('w', encoding='utf-8-sig', newline='') as handle:
                    writer = csv.writer(handle)
                    # CSV is analysis data, not an Excel UI; values stay exact, including formula-like source text.
                    writer.writerow(all_headers)
                    writer.writerows(rows)
    if issues:
        sheet(name_for(tr('Feil og merknader','Errors and notes')),
              [tr('Dokument','Document'), tr('Kjøring','Run'), tr('Forsøk','Attempt'), tr('Type','Type'), tr('Melding','Message')],
              issues, {1:30, 2:16, 3:16, 4:24, 5:110})
    if any(len(record['detail']['forsok']) > 1 for record in records):
        attempts = []
        for record in records:
            run = record['run']
            for item in record['detail']['forsok']:
                attempt = item['forsok']
                validation = item.get('result_validation') or {}
                attempts.append([record['document']['navn'], run['id'], attempt['id'], attempt['nr'],
                                 tr('Gjeldende', 'Current') if attempt['id'] == run.get('gjeldende_forsok_id') else
                                 tr('Tidligere', 'Earlier'), attempt['status'],
                                 tr('Bestått', 'Passed') if validation.get('gyldig') else
                                 tr('Ikke bestått', 'Not passed') if validation else tr('Ikke kontrollert', 'Not checked'),
                                 item.get('review_status', ''), attempt.get('feil') or
                                 '\n'.join(e.get('melding', '') for e in validation.get('feil', [])),
                                 attempt.get('startet'), attempt.get('avsluttet')])
        sheet(name_for(tr('Forsøk', 'Attempts')),
              [tr('Dokument', 'Document'), tr('Kjøring', 'Run'), tr('Forsøk', 'Attempt'),
               tr('Nummer', 'Number'), tr('Rolle', 'Role'), tr('Status', 'Status'),
               tr('Validering', 'Validation'), tr('Menneskelig kontroll', 'Human review'),
               tr('Feil', 'Error'), tr('Startet', 'Started'), tr('Avsluttet', 'Finished')],
              attempts, {1:30, 2:16, 3:18, 5:16, 8:22, 9:80, 10:24, 11:24})
    sheet(dictionary_name, [tr('Dataark','Data sheet'), tr('Kolonne','Column'), tr('Variabelsti','Variable path'),
                           tr('Forklaring','Definition'), tr('Datatype','Data type'), tr('Planversjon','Plan version'),
                           tr('Tabellsti','Table path'), tr('Én rad betyr','Row unit')], dictionary,
          {1:28, 2:30, 3:35, 4:90, 5:25, 6:20, 7:35, 8:75})
    if overflow:
        sheet(texts_name, [tr('Celle','Cell'), tr('Del','Part'), tr('Tekst','Text')], overflow, {1:35, 2:10, 3:120})
    filename = tr('Resultater.xlsx', 'Results.xlsx')
    book.save(directory / filename)
    book.close()
    return filename, [{'sheet': results_name, 'rows': len(records), 'row_unit': 'One run' if row_scope == 'runs' else 'One document'}] + [
        {'sheet': n, 'plan_version_id': p, 'rows': c, 'row_unit': u} for n,p,c,u in summaries]
