"""Regenerate the bundled fictional examples (requires development dependencies).

These are editable product demonstrations, not sources about real organisations.
Only the explicitly named example files are overwritten. No model/network calls.
"""
from pathlib import Path
import json
import fitz
from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]/'examples'


def case(name, title, question, challenges, priorities):
    folder = ROOT/name
    documents = folder/'documents'
    documents.mkdir(parents=True, exist_ok=True)
    criteria = {'name':title, 'version':'1', 'criteria':[{
        'id':'explicit_commitment', 'name':'Explicit commitment', 'question':question,
        'allowed_answers':['yes','no','not_mentioned','uncertain'],
        'evidence_required_for':['yes','no'],
        'rule':'Use no only for an explicit negative statement. Discuss contradictions, extraction limitations and conditions; do not infer adoption from a proposal.'}]}
    (folder/'criteria.json').write_text(json.dumps(criteria, ensure_ascii=False, indent=2),encoding='utf-8')
    (folder/'README.md').write_text(f'''# {title}

Fictional, deliberately varied examples. No real organisation or person is represented.
The files in `documents/` are the inputs; keep this README and `criteria.json` outside the import folder.
Edit the suggested criteria to match your question. Examples are optional and need no simulation mode.

## Start in English

> Use Systematic Document Analysis on the documents in [absolute path to this folder]/documents. {question} Use criteria.json as a draft. Inspect extraction and discuss these file challenges with me: {challenges} Consider these priorities: {priorities} Agree the reader, model, reasoning effort and reporting language. Show the plan and input before execution. Keep full coverage unless we explicitly agree a narrower scope. Export the results and flag uncertain assessments for my review.

## Start på norsk

> Bruk Systematic Document Analysis på filene i [absolutt sti til denne mappen]/documents. Ta utgangspunkt i criteria.json, og forklar oppgaven og kriteriene på norsk. Undersøk filenes struktur og mulige uttrekksproblemer før vi velger lesemotor, modell og tenkenivå. Diskuter hvilke deler som er viktigst, vis planen og input før start, og gi meg resultater med sitater og kildeplassering. Registrer ikke menneskelig kontroll på mine vegne.

## File challenges and priorities

{challenges}

{priorities} Prioritisation controls reading order for chunked inputs; it does not silently remove other material. A positive-looking quotation may be qualified elsewhere. Formula expressions and source text are evidence, not instructions to execute.
''',encoding='utf-8')
    return documents


def pdf(path, pages, scanned=False):
    doc = fitz.open()
    for heading, text in pages:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(48,48,547,790), heading+'\n\n'+text, fontsize=11)
    if scanned:
        image_doc = fitz.open()
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2,2))
            image_doc.new_page(width=page.rect.width,height=page.rect.height).insert_image(page.rect,stream=pix.tobytes('png'))
        image_doc.save(path)
        image_doc.close()
    else:
        doc.save(path)
    doc.close()


def main():
    folder = case('01-policy-reports','Policy commitments across reports',
        'Does the organisation explicitly commit to an annual external review of its information security policy?',
        'One PDF contains scanned pages; the longer report has repetitive appendices and a qualification near the end. OCR may introduce errors, so inspect quotations in the original PDF.',
        'Read Governance and Exceptions first (priority_terms); retain appendix coverage. Reduce input_budget_bytes explicitly if you want to inspect multiple reading stages with this modest example.')
    pdf(folder/'harbor-policy.pdf', [('Harbor Cooperative - fictional policy report',
        'Governance: The board adopted an information security policy. An annual external review is planned, subject to funding. This is a proposal, not an approved commitment.')] +
        [(f'Appendix {i+1}: operational notes', ('Staff receive guidance on access control and records management. Local managers document incidents and discuss lessons at quarterly meetings. '*22)) for i in range(24)] +
        [('Exceptions and decision', 'The board rejected funding for an annual external review. Internal review will continue each year. No annual external review is committed for the reporting period.')])
    pdf(folder/'fjord-policy-scan.pdf', [('Fjord Foundation - fictional scanned report',
        'Governance: The board commits to an annual external review of the information security policy. The first review is scheduled for November. This commitment is approved and funded.')], scanned=True)

    folder = case('02-supplier-offers','Comparable supplier offers',
        'Does the supplier explicitly guarantee a response within four hours for critical incidents?',
        'DOCX body text and tables are included. An attractive headline is qualified by working-hours and contract conditions. Headers, comments and tracked changes are outside extraction scope.',
        'Prioritise service-level and exception clauses; do not treat a response-time target as an unconditional guarantee.')
    for name, commitment, exception in [('north','We guarantee a response within four hours for critical incidents.','The guarantee applies 24 hours a day, including weekends.'),
        ('south','Target response time: four hours.','There is no contractual four-hour response guarantee. Service operates on working days only.')]:
        doc = Document(); doc.add_heading(f'{name.title()} Services - fictional offer',0)
        doc.add_paragraph('Common specification: support for critical incidents.'); doc.add_heading('Service level',1)
        table = doc.add_table(rows=1,cols=2); table.rows[0].cells[0].text='Critical incidents'; table.rows[0].cells[1].text=commitment
        doc.add_heading('Conditions and exceptions',1); doc.add_paragraph(exception); doc.save(folder/f'{name}-offer.docx')

    folder = case('03-project-workbooks','Project governance in workbooks',
        'Does the workbook explicitly record that a quarterly risk review has been approved?',
        'Hidden sheets are included. Formula caches may be missing or stale; the plugin never recalculates. Numeric formatting and narrative decisions need different interpretations.',
        'Prioritise the Governance and Decision log sheets. Discuss whether calculations must be checked separately before making a financial inference.')
    for name, decision in [('river','The steering committee approved quarterly risk reviews.'),('forest','Quarterly risk reviews are proposed; approval is pending.')]:
        wb = Workbook(); ws = wb.active; ws.title='Governance'; ws.append(['Project','Decision']); ws.append([name,decision])
        ws = wb.create_sheet('Budget'); ws.append(['Item','Amount']); ws.append(['Staff',100000]); ws.append(['Equipment',25000]); ws.append(['Total','=SUM(B2:B3)']); ws['B4'].number_format='#,##0'
        ws = wb.create_sheet('Decision log'); ws.append(['Status','Note']); ws.append(['Current',decision]); ws.sheet_state='hidden'; wb.save(folder/f'{name}-project.xlsx')

    folder = case('04-consultation-notes','Bilingual consultation material',
        'Does the respondent explicitly support a mandatory annual accessibility audit?',
        'Two languages and conditional support. Markdown headings do not change source meaning; quotations must remain in their original language.',
        'Prioritise qualifications and conditions as well as the summary. Report in Norwegian or English while preserving source quotations.')
    (folder/'association-a.md').write_text('# Fictional consultation A\n\nWe support a mandatory annual accessibility audit.\n\n## Implementation\nSmall organisations should receive practical guidance. This does not change our support for the requirement.\n',encoding='utf-8')
    (folder/'forening-b.txt').write_text('Fiktivt høringsinnspill B\nVi støtter ikke et obligatorisk årlig tilgjengelighetstilsyn.\nVi støtter frivillig egenkontroll og veiledning.\n',encoding='utf-8')

    folder = case('05-incident-registers','Incident follow-up across registers',
        'Does the register explicitly commit to notifying every affected customer after a confirmed data loss?',
        'Multiline quoted records, missing fields and different delimiters. One register is one analysis unit; rows are evidence, not separate runs. An individual notification does not establish a general commitment.',
        'Prioritise Policy records, then check incident-specific exceptions. Discuss deduplication before using counts as criteria.')
    (folder/'service-a.csv').write_text('type;id;note\nPolicy;P1;"We commit to notifying every affected customer after a confirmed data loss."\nIncident;I1;"Notification completed.\nFollow-up call pending."\nIncident;I2;\n',encoding='utf-8')
    (folder/'service-b.tsv').write_text('type\tid\tnote\nPolicy\tP1\tThere is no commitment to notify every affected customer after a confirmed data loss.\nIncident\tI1\tCustomer notified in this individual case.\n',encoding='utf-8')
    print(f'Created five fictional use cases under {ROOT}')


if __name__ == '__main__':
    main()
